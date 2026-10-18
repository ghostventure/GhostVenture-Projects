from pathlib import Path
import json, re, math, zipfile, xml.etree.ElementTree as ET
import pymupdf as fitz
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from PIL import Image, ImageOps, ImageDraw
ROOT=Path(__file__).resolve().parent
model=json.loads((ROOT/'financial-model.json').read_text())
path=ROOT/'Overhead-LLC-Financial-Model.xlsx'
wb=load_workbook(path,data_only=False)
cache={}
def flatten(v):
    if isinstance(v,(tuple,list)):
        return [x for y in v for x in flatten(y)]
    return [v]
def value(sheet,ref):
    key=(sheet,ref.replace('$',''))
    if key in cache:return cache[key]
    cell=wb[key[0]][key[1]]; raw=cell.value
    if not isinstance(raw,str) or not raw.startswith('='):return raw or 0
    variables={}
    def reference(match):
        target=match.group(1) or sheet
        first=match.group(2).replace('$','');last=match.group(3)
        if last:
            lo,top,hi,bottom=range_boundaries(first+':'+last.replace('$',''))
            result=[[value(target,wb[target].cell(r,c).coordinate) for c in range(lo,hi+1)] for r in range(top,bottom+1)]
        else:result=value(target,first)
        name='_v'+str(len(variables));variables[name]=result;return name
    expression=re.sub(r'(?:(\w+)!)?(\$?[A-Z]{1,3}\$?\d+)(?::(\$?[A-Z]{1,3}\$?\d+))?',reference,raw[1:])
    functions={'SUM':lambda *v:sum(flatten(v)),'SUMPRODUCT':lambda a,b:sum(x*y for x,y in zip(flatten(a),flatten(b))),'MAX':lambda *v:max(flatten(v)),'INT':math.floor,'INDEX':lambda a,r,c:a[int(r)-1][int(c)-1]}
    result=eval(expression,{'__builtins__':{}},dict(functions,**variables));cache[key]=result;return result
keys=['month','year','opening','churned','adds','ending','billed','gross_sales','refunds','revenue','fees','delivery','onboarding','cogs','gross_profit','opex','launch','operating_profit','tax','net_income','ending_cash','mrr','arr']
checked=0
for name,case in model['cases'].items():
    for rownum,record in enumerate(case['rows'],2):
        for colnum,key in enumerate(keys,1):
            actual=value(name,wb[name].cell(rownum,colnum).coordinate)
            assert abs(actual-record[key])<1e-6,(name,rownum,key,actual,record[key])
            checked+=1
for sheet in wb:
    for row in sheet:
        for cell in row:
            if cell.data_type=='f':value(sheet.title,cell.coordinate)
# Store verified formula results as preview caches. Excel still recalculates after edits.
ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
temp=path.with_suffix('.verified.xlsx')
with zipfile.ZipFile(path) as source,zipfile.ZipFile(temp,'w',zipfile.ZIP_DEFLATED) as out:
    for item in source.infolist():
        data=source.read(item.filename)
        match=re.fullmatch(r'xl/worksheets/sheet(\d+)\.xml',item.filename)
        if match:
            sheet=wb.worksheets[int(match.group(1))-1];tree=ET.fromstring(data)
            for c in tree.findall('.//s:c',ns):
                key=(sheet.title,c.attrib['r'])
                if c.find('s:f',ns) is not None and key in cache:
                    v=c.find('s:v',ns)
                    if v is None:v=ET.SubElement(c,'{'+ns['s']+'}v')
                    v.text=str(cache[key])
            data=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
        out.writestr(item,data)
temp.replace(path)
cached=load_workbook(path,data_only=True)
assert abs(cached['Base']['U37'].value-model['cases']['Base']['annual'][2]['cash'])<1e-6

pdfpath=ROOT/'Overhead-LLC-Business-Plan.pdf';doc=fitz.open(pdfpath)
assert len(doc)==27,len(doc)
metadata=doc.metadata;metadata.update(title='Overhead LLC — Business Plan',author='Overhead LLC',subject='General-purpose business plan with research and illustrative financial projections',keywords='Overhead LLC, OverHead, business plan, financial model, September 2026')
doc.set_metadata(metadata)
temp_pdf=ROOT/'Overhead-LLC-Business-Plan-verified.pdf';doc.save(temp_pdf,garbage=4,deflate=True);doc.close();temp_pdf.replace(pdfpath)
doc=fitz.open(pdfpath)
report=[];thumbs=[]
for i,p in enumerate(doc,1):
    text=p.get_text()
    assert 'Overhead LLC' in text or 'OVERHEAD LLC' in text,(i,'Company name')
    assert f'{i:02} / 27' in text,(i,'Footer missing')
    assert '\ufffd' not in text,(i,'Replacement character')
    assert len(text)>350,(i,'Short or blank page')
    for b in p.get_text('blocks'):
        assert b[0]>=-1 and b[1]>=-1 and b[2]<=p.rect.width+1 and b[3]<=p.rect.height+1,(i,'Text outside page')
    for link in p.get_links():
        if link['kind']==fitz.LINK_GOTO:assert 0<=link['page']<27
    report.append({'page':i,'words':len(text.split()),'links':len(p.get_links()),'dimensions':[p.rect.width,p.rect.height]})
    pix=p.get_pixmap(matrix=fitz.Matrix(.42,.42),alpha=False)
    img=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
    tile=Image.new('RGB',(278,360),'#e4eaee');tile.paste(img,((278-img.width)//2,8));ImageDraw.Draw(tile).text((12,343),f'Page {i}',fill='#183448');thumbs.append(tile)
    if i in [1,8,18,19,20,21,27]:p.get_pixmap(matrix=fitz.Matrix(1.4,1.4),alpha=False).save(ROOT/'assets'/f'pdf-page-{i}.png')
for start in [0,15]:
    batch=thumbs[start:start+15];sheet=Image.new('RGB',(278*5,360*math.ceil(len(batch)/5)),'#e4eaee')
    for j,img in enumerate(batch):sheet.paste(img,((j%5)*278,(j//5)*360))
    sheet.save(ROOT/'assets'/f'contact-sheet-{start+1}.png')
layout=json.loads((ROOT/'layout-audit.json').read_text());assert not any(x['overflow'] or x['wide'] for x in layout)
summary={'pages':len(doc),'words':sum(x['words'] for x in report),'clickable_links':sum(x['links'] for x in report),'formula_outputs_checked':checked,'formula_cache_count':len(cache),'layout_overflows':0,'pdf_text_bounds':'PASS','financial_reconciliation':'PASS','workbook_vs_independent_model':'PASS','page_details':report}
(ROOT/'verification.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in summary.items() if k!='page_details'},indent=2))
