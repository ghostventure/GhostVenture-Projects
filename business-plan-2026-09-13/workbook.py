from pathlib import Path
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as col
from openpyxl.workbook.properties import CalcProperties
ROOT=Path(__file__).resolve().parent
model=json.loads((ROOT/'financial-model.json').read_text())
wb=Workbook(); info=wb.active;info.title='Read me'
notes=[['Overhead LLC | 36-month planning model'],['General-purpose business plan | September 13, 2026'],['All figures are illustrative assumptions, not actual results.'],['Change blue input values on Assumptions and Acquisition. Formula sheets recalculate in Excel.'],['Month 1 starts with funded execution; months 1-2 have no paid additions.'],['Base, Downside, and Upside use the same operating budget and $175,000 capital scenario.'],['Monthly churn affects opening accounts. New additions contribute half a billed month.'],['Net revenue deducts a 1% refund/credit allowance. Fees use gross charges.'],['No debt, capex, working capital lag, historical assets, or existing balances are included.'],['The 25% annual tax reserve is a planning assumption, not an LLC tax determination.'],['Model values were independently calculated in Python and exported to the companion CSV/JSON.'],['Formula cells may show blank previews until opened in Excel or another recalculating spreadsheet app.']]
for r in notes:info.append(r)
info.column_dimensions['A'].width=112
a=wb.create_sheet('Assumptions')
inputs=[['Input','Value','Basis'],['Opening capital',175000,'Illustrative new contribution; not existing company cash'],['Silver monthly price',59,'Published'],['Gold monthly price',129,'Published'],['Black monthly price',299,'Published'],['Silver mix',.4,'Assumption'],['Gold mix',.45,'Assumption'],['Black mix',.15,'Assumption'],['Blended ARPA','=SUMPRODUCT(B3:B5,B6:B8)','Calculated'],['Base monthly churn',.02,'Assumption'],['Refunds/credits rate',.01,'Assumption'],['Processing + Billing rate',.036,'2.9% domestic card + 0.7% Stripe Billing'],['Fixed payment fee',.30,'Per billed workspace-month approximation'],['Infrastructure / billable month',5,'Assumption'],['Support / billable month',10,'Assumption'],['Onboarding / new paid workspace',60,'Assumption'],['Tax reserve',.25,'Applied to positive annual operating income; paid at year-end'],['Minimum cash reserve',30000,'Management target'],['Downside acquisition factor',.6,'Assumption'],['Downside churn',.035,'Assumption'],['Downside price factor',.9,'Assumption'],['Upside acquisition factor',1.4,'Assumption'],['Upside churn',.015,'Assumption']]
for r in inputs:a.append(r)
a.append([]);a.append(['Annual budget','Year 1','Year 2','Year 3'])
for label,key in [('Product','product'),('Founder','founder'),('Marketing','marketing'),('General','general')]:a.append([label]+[x[key]*12 for x in model['opex']])
a.append(['Monthly fixed opex','=SUM(B26:B29)/12','=SUM(C26:C29)/12','=SUM(D26:D29)/12'])
acq=wb.create_sheet('Acquisition');acq.append(['Planning month','Base gross new paid','Launch expense'])
for i,n in enumerate(model['base_adds'],1):acq.append([i,n,10000 if i==1 else 5000 if i==2 else 0])
headers=['Month','Year','Opening paid','Churned','New paid','Ending paid','Billed equivalent','Gross charges','Credits/refunds','Net revenue','Payment fees','Service delivery','Onboarding','Total COGS','Gross profit','Fixed opex','Launch costs','Operating profit','Tax reserve','Net income','Ending cash','Exit MRR','Exit ARR']
for name in ['Base','Downside','Upside']:
 s=wb.create_sheet(name);s.append(headers)
 factor='1' if name=='Base' else 'Assumptions!$B$19' if name=='Downside' else 'Assumptions!$B$22'
 churn='Assumptions!$B$10' if name=='Base' else 'Assumptions!$B$20' if name=='Downside' else 'Assumptions!$B$23'
 arpa='Assumptions!$B$9'+('*Assumptions!$B$21' if name=='Downside' else '')
 for month in range(1,37):
  r=month+1;start=r-11;prev='0' if month==1 else f'F{r-1}'
  formulas=[month,f'=INT((A{r}-1)/12)+1',f'={prev}',f'=C{r}*{churn}',f'=Acquisition!B{r}*{factor}',f'=C{r}-D{r}+E{r}',f'=C{r}-D{r}+E{r}/2',f'=G{r}*{arpa}',f'=H{r}*Assumptions!$B$11',f'=H{r}-I{r}',f'=H{r}*Assumptions!$B$12+G{r}*Assumptions!$B$13',f'=G{r}*(Assumptions!$B$14+Assumptions!$B$15)',f'=E{r}*Assumptions!$B$16',f'=SUM(K{r}:M{r})',f'=J{r}-N{r}',f'=INDEX(Assumptions!$B$30:$D$30,1,B{r})',f'=Acquisition!C{r}',f'=O{r}-P{r}-Q{r}',f'=MAX(SUM(R{start}:R{r}),0)*Assumptions!$B$17' if month%12==0 else '=0',f'=R{r}-S{r}',f'='+('Assumptions!$B$2' if month==1 else f'U{r-1}')+f'+T{r}',f'=F{r}*{arpa}',f'=V{r}*12']
  s.append(formulas)
 for row in s.iter_rows(min_row=2):
  for cell in row[2:]:cell.number_format='#,##0.00;(#,##0.00)'
 s.freeze_panes='C2';s.auto_filter.ref=s.dimensions
 s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
summary=wb.create_sheet('Annual summary');summary.append(['Case','Year','Ending paid','Net revenue','Gross profit','Operating profit','Net income','Ending cash','Exit ARR'])
for name in ['Base','Downside','Upside']:
 for year in range(1,4):
  start=2+(year-1)*12;end=start+11
  summary.append([name,year,f'={name}!F{end}',f'=SUM({name}!J{start}:J{end})',f'=SUM({name}!O{start}:O{end})',f'=SUM({name}!R{start}:R{end})',f'=SUM({name}!T{start}:T{end})',f'={name}!U{end}',f'={name}!W{end}'])
for row in summary.iter_rows(min_row=2):
 for c in row[2:]:c.number_format='#,##0.00;(#,##0.00)'
for sheet in wb:
 for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='143246');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True)
 sheet.row_dimensions[1].height=32
 if sheet.title!='Read me':
  for column in range(1,sheet.max_column+1):sheet.column_dimensions[col(column)].width=19
 for row in sheet.iter_rows(min_row=2):
  for c in row:
   if c.data_type=='f':c.font=Font(color='173042')
   elif isinstance(c.value,(int,float)):c.font=Font(color='1675BB')
 sheet.freeze_panes='B2'
a.column_dimensions['A'].width=35;a.column_dimensions['C'].width=65
for r in [6,7,8,10,11,12,17,19,20,21,22,23]:a.cell(r,2).number_format='0.0%'
wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
wb.save(ROOT/'Overhead-LLC-Financial-Model.xlsx')
assert a['B30'].value=='=SUM(B26:B29)/12'
print('Workbook saved: 36 months, 3 scenarios, formula-driven assumptions.')
